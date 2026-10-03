# ============================================================
# INTELLIFEET
# SERVICIO - CONTROL DIARIO
# ============================================================

from datetime import datetime, date
from zoneinfo import ZoneInfo
from pathlib import Path
from io import BytesIO
import os
import uuid

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter

from werkzeug.utils import secure_filename

from extensions import db
from models import (
    ControlDiario,
    Usuario,
    Vehiculo,
    Maquinaria,
    ActivoOperador
)


# ============================================================
# CONFIGURACIÓN
# ============================================================

BOGOTA = ZoneInfo("America/Bogota")

UPLOAD_FOLDER = os.path.join(
    "uploads",
    "controles_diarios"
)

ALLOWED_EXTENSIONS = {
    "jpg",
    "jpeg",
    "png",
    "webp",
    "pdf"
}


# ============================================================
# UTILIDADES
# ============================================================

def archivo_permitido(nombre):
    """
    Verifica que el archivo tenga una extensión permitida.
    """

    if not nombre:
        return False

    if "." not in nombre:
        return False

    extension = nombre.rsplit(
        ".",
        1
    )[1].lower()

    return extension in ALLOWED_EXTENSIONS


def obtener_extension(nombre):
    """
    Obtiene la extensión del archivo.
    """

    if not nombre or "." not in nombre:
        return ""

    return nombre.rsplit(
        ".",
        1
    )[1].lower()


def obtener_usuario(usuario_id):
    """
    Obtiene el usuario actual.
    """

    try:
        usuario_id = int(usuario_id)
    except (TypeError, ValueError):
        return None

    return Usuario.query.get(usuario_id)


def validar_rol_admin(usuario):
    """
    Determina si el usuario puede administrar
    los controles diarios.
    """

    if not usuario or not usuario.rol:
        return False

    rol = usuario.rol.nombre.lower().strip()

    return rol in [
        "admin",
        "administrador",
        "supervisor"
    ]


# ============================================================
# ACTIVOS DEL OPERADOR
# ============================================================

def obtener_activos_operador(usuario_id):
    """
    Obtiene únicamente los vehículos y maquinarias
    actualmente asignados al operador.
    """

    asignaciones = ActivoOperador.query.filter_by(
        usuario_id=usuario_id,
        activo=True
    ).all()

    vehiculos = []
    maquinarias = []

    for asignacion in asignaciones:

        # ----------------------------------------------------
        # VEHÍCULO
        # ----------------------------------------------------

        if asignacion.vehiculo:

            vehiculo = asignacion.vehiculo

            vehiculos.append({
                "id": vehiculo.id,
                "placa": vehiculo.placa,
                "tipo": (
                    vehiculo.tipo_vehiculo.nombre
                    if vehiculo.tipo_vehiculo
                    else None
                ),
                "marca": vehiculo.marca,
                "modelo": vehiculo.modelo,
                "estado": vehiculo.estado
            })

        # ----------------------------------------------------
        # MAQUINARIA
        # ----------------------------------------------------

        if asignacion.maquinaria:

            maquinaria = asignacion.maquinaria

            maquinarias.append({
                "id": maquinaria.id,
                "codigo": maquinaria.codigo,
                "tipo": (
                    maquinaria.tipo_maquinaria.nombre
                    if maquinaria.tipo_maquinaria
                    else None
                ),
                "marca": maquinaria.marca,
                "modelo": maquinaria.modelo,
                "horometro_actual": (
                    maquinaria.horometro_actual
                ),
                "estado": maquinaria.estado
            })

    return {
        "vehiculos": vehiculos,
        "maquinarias": maquinarias
    }


# ============================================================
# VALIDAR ASIGNACIÓN DEL ACTIVO
# ============================================================

def validar_activo_operador(
    usuario_id,
    vehiculo_id=None,
    maquinaria_id=None
):
    """
    Verifica que el activo pertenezca a una asignación
    activa del operador.
    """

    if not vehiculo_id and not maquinaria_id:
        return None

    if vehiculo_id and maquinaria_id:
        return None

    query = ActivoOperador.query.filter_by(
        usuario_id=usuario_id,
        activo=True
    )

    if vehiculo_id:

        query = query.filter_by(
            vehiculo_id=vehiculo_id
        )

    elif maquinaria_id:

        query = query.filter_by(
            maquinaria_id=maquinaria_id
        )

    return query.first()


# ============================================================
# GUARDAR ARCHIVO
# ============================================================

def guardar_archivo(
    archivo,
    fecha_control
):
    """
    Guarda el archivo organizado por:

    uploads/
        controles_diarios/
            año/
                mes/
                    archivo
    """

    if not archivo:
        raise ValueError(
            "Debe adjuntar una evidencia"
        )

    if not archivo.filename:
        raise ValueError(
            "El archivo no tiene nombre"
        )

    if not archivo_permitido(
        archivo.filename
    ):
        raise ValueError(
            "Tipo de archivo no permitido. "
            "Use JPG, JPEG, PNG, WEBP o PDF."
        )

    # --------------------------------------------------------
    # CARPETA
    # --------------------------------------------------------

    carpeta = os.path.join(
        UPLOAD_FOLDER,
        str(fecha_control.year),
        f"{fecha_control.month:02d}"
    )

    os.makedirs(
        carpeta,
        exist_ok=True
    )

    # --------------------------------------------------------
    # NOMBRE ORIGINAL
    # --------------------------------------------------------

    nombre_original = secure_filename(
        archivo.filename
    )

    if not nombre_original:
        raise ValueError(
            "Nombre de archivo inválido"
        )

    extension = obtener_extension(
        nombre_original
    )

    # --------------------------------------------------------
    # NOMBRE ÚNICO
    # --------------------------------------------------------

    nombre_archivo = (
        f"{uuid.uuid4().hex}.{extension}"
    )

    # --------------------------------------------------------
    # RUTA FÍSICA
    # --------------------------------------------------------

    ruta_fisica = os.path.join(
        carpeta,
        nombre_archivo
    )

    archivo.save(
        ruta_fisica
    )

    # --------------------------------------------------------
    # RUTA PARA BD / URL
    # --------------------------------------------------------

    archivo_path = os.path.join(
        "controles_diarios",
        str(fecha_control.year),
        f"{fecha_control.month:02d}",
        nombre_archivo
    ).replace("\\", "/")

    return {
        "archivo_path": archivo_path,
        "archivo_nombre": nombre_original,
        "archivo_tipo": extension,
        "ruta_fisica": ruta_fisica
    }


# ============================================================
# CREAR CONTROL
# ============================================================

def crear_control_diario(
    usuario_id,
    fecha_control,
    vehiculo_id=None,
    maquinaria_id=None,
    archivo=None,
    observaciones=None
):
    """
    Crea un nuevo control diario.
    """

    usuario = obtener_usuario(
        usuario_id
    )

    if not usuario:
        raise ValueError(
            "Usuario no encontrado"
        )

    # --------------------------------------------------------
    # FECHA
    # --------------------------------------------------------

    if isinstance(
        fecha_control,
        str
    ):

        try:

            fecha_control = date.fromisoformat(
                fecha_control
            )

        except ValueError:

            raise ValueError(
                "La fecha no tiene un formato válido"
            )

    if not isinstance(
        fecha_control,
        date
    ):

        raise ValueError(
            "Fecha inválida"
        )

    # --------------------------------------------------------
    # ACTIVO
    # --------------------------------------------------------

    if not vehiculo_id and not maquinaria_id:

        raise ValueError(
            "Debe seleccionar un vehículo "
            "o una maquinaria"
        )

    if vehiculo_id and maquinaria_id:

        raise ValueError(
            "No puede seleccionar vehículo "
            "y maquinaria al mismo tiempo"
        )

    # --------------------------------------------------------
    # CONVERTIR IDS
    # --------------------------------------------------------

    try:

        if vehiculo_id is not None:
            vehiculo_id = int(
                vehiculo_id
            )

        if maquinaria_id is not None:
            maquinaria_id = int(
                maquinaria_id
            )

    except (TypeError, ValueError):

        raise ValueError(
            "El ID del activo no es válido"
        )

    # --------------------------------------------------------
    # VALIDAR ASIGNACIÓN
    # --------------------------------------------------------

    asignacion = validar_activo_operador(
        usuario_id=usuario.id,
        vehiculo_id=vehiculo_id,
        maquinaria_id=maquinaria_id
    )

    if not asignacion:

        raise PermissionError(
            "El activo seleccionado "
            "no está asignado al operador"
        )

    # --------------------------------------------------------
    # GUARDAR ARCHIVO
    # --------------------------------------------------------

    archivo_info = guardar_archivo(
        archivo=archivo,
        fecha_control=fecha_control
    )

    # --------------------------------------------------------
    # CREAR REGISTRO
    # --------------------------------------------------------

    control = ControlDiario(

        fecha=fecha_control,

        usuario_id=usuario.id,

        vehiculo_id=vehiculo_id,

        maquinaria_id=maquinaria_id,

        archivo_path=(
            archivo_info["archivo_path"]
        ),

        archivo_nombre=(
            archivo_info["archivo_nombre"]
        ),

        archivo_tipo=(
            archivo_info["archivo_tipo"]
        ),

        estado="PENDIENTE",

        observaciones=observaciones
    )

    try:

        db.session.add(
            control
        )

        db.session.commit()

    except Exception:

        db.session.rollback()

        # ----------------------------------------------------
        # Si BD falla, intentamos eliminar el archivo
        # para no dejar basura en uploads.
        # ----------------------------------------------------

        try:

            if os.path.exists(
                archivo_info["ruta_fisica"]
            ):

                os.remove(
                    archivo_info["ruta_fisica"]
                )

        except Exception:
            pass

        raise

    return control


# ============================================================
# OBTENER CONTROL
# ============================================================

def obtener_control_diario(
    control_id
):
    return ControlDiario.query.get(
        control_id
    )


# ============================================================
# CONTROLES DEL OPERADOR
# ============================================================

def obtener_controles_operador(
    usuario_id,
    fecha_desde=None,
    fecha_hasta=None
):
    """
    Obtiene únicamente los controles creados
    por el operador.
    """

    query = ControlDiario.query.filter_by(
        usuario_id=usuario_id
    )

    # --------------------------------------------------------
    # FECHA DESDE
    # --------------------------------------------------------

    if fecha_desde:

        if isinstance(
            fecha_desde,
            str
        ):

            fecha_desde = date.fromisoformat(
                fecha_desde
            )

        query = query.filter(
            ControlDiario.fecha >= fecha_desde
        )

    # --------------------------------------------------------
    # FECHA HASTA
    # --------------------------------------------------------

    if fecha_hasta:

        if isinstance(
            fecha_hasta,
            str
        ):

            fecha_hasta = date.fromisoformat(
                fecha_hasta
            )

        query = query.filter(
            ControlDiario.fecha <= fecha_hasta
        )

    return query.order_by(
        ControlDiario.fecha.desc(),
        ControlDiario.id.desc()
    ).all()


# ============================================================
# FILTRO ADMINISTRATIVO
# ============================================================

# ============================================================
# FILTRO ADMINISTRATIVO
# ============================================================

def obtener_controles_admin(
    mes=None,
    anio=None,
    vehiculo_id=None,
    maquinaria_id=None,
    usuario_id=None,
    estado=None,
    fecha_desde=None,
    fecha_hasta=None
):
    """
    Consulta administrativa con filtros.

    Puede utilizarse tanto para el listado administrativo
    como para la generación del Excel consolidado.
    """

    query = ControlDiario.query

    # --------------------------------------------------------
    # CONVERTIR MES / AÑO
    # --------------------------------------------------------

    if mes:
        try:
            mes = int(mes)
        except (TypeError, ValueError):
            raise ValueError("Mes inválido")

        if mes < 1 or mes > 12:
            raise ValueError("Mes inválido")

    if anio:
        try:
            anio = int(anio)
        except (TypeError, ValueError):
            raise ValueError("Año inválido")

        if anio < 2000 or anio > 2100:
            raise ValueError("Año inválido")

    # --------------------------------------------------------
    # MES + AÑO
    # --------------------------------------------------------

    if mes and anio:

        fecha_inicio = date(
            anio,
            mes,
            1
        )

        if mes == 12:

            fecha_fin = date(
                anio + 1,
                1,
                1
            )

        else:

            fecha_fin = date(
                anio,
                mes + 1,
                1
            )

        query = query.filter(
            ControlDiario.fecha >= fecha_inicio,
            ControlDiario.fecha < fecha_fin
        )

    # --------------------------------------------------------
    # SOLAMENTE AÑO
    # --------------------------------------------------------

    elif anio:

        fecha_inicio = date(
            anio,
            1,
            1
        )

        fecha_fin = date(
            anio + 1,
            1,
            1
        )

        query = query.filter(
            ControlDiario.fecha >= fecha_inicio,
            ControlDiario.fecha < fecha_fin
        )

    # --------------------------------------------------------
    # SOLAMENTE MES
    # --------------------------------------------------------

    elif mes:

        from sqlalchemy import extract

        query = query.filter(
            extract(
                "month",
                ControlDiario.fecha
            ) == mes
        )

    # --------------------------------------------------------
    # FECHA DESDE
    # --------------------------------------------------------

    if fecha_desde:

        try:

            if isinstance(
                fecha_desde,
                str
            ):

                fecha_desde = date.fromisoformat(
                    fecha_desde
                )

        except ValueError:

            raise ValueError(
                "Fecha desde inválida"
            )

        query = query.filter(
            ControlDiario.fecha >= fecha_desde
        )

    # --------------------------------------------------------
    # FECHA HASTA
    # --------------------------------------------------------

    if fecha_hasta:

        try:

            if isinstance(
                fecha_hasta,
                str
            ):

                fecha_hasta = date.fromisoformat(
                    fecha_hasta
                )

        except ValueError:

            raise ValueError(
                "Fecha hasta inválida"
            )

        query = query.filter(
            ControlDiario.fecha <= fecha_hasta
        )

    # --------------------------------------------------------
    # VEHÍCULO
    # --------------------------------------------------------

    if vehiculo_id:

        try:
            vehiculo_id = int(vehiculo_id)
        except (TypeError, ValueError):
            raise ValueError(
                "ID de vehículo inválido"
            )

        query = query.filter(
            ControlDiario.vehiculo_id == vehiculo_id
        )

    # --------------------------------------------------------
    # MAQUINARIA
    # --------------------------------------------------------

    if maquinaria_id:

        try:
            maquinaria_id = int(maquinaria_id)
        except (TypeError, ValueError):
            raise ValueError(
                "ID de maquinaria inválido"
            )

        query = query.filter(
            ControlDiario.maquinaria_id == maquinaria_id
        )

    # --------------------------------------------------------
    # OPERADOR
    # --------------------------------------------------------

    if usuario_id:

        try:
            usuario_id = int(usuario_id)
        except (TypeError, ValueError):
            raise ValueError(
                "ID de operador inválido"
            )

        query = query.filter(
            ControlDiario.usuario_id == usuario_id
        )

    # --------------------------------------------------------
    # ESTADO
    # --------------------------------------------------------

    if estado:

        query = query.filter(
            ControlDiario.estado == estado.upper().strip()
        )

    # --------------------------------------------------------
    # RESULTADO
    # --------------------------------------------------------

    return query.order_by(
        ControlDiario.fecha.desc(),
        ControlDiario.id.desc()
    ).all()
    
# ============================================================
# VALIDAR CONTROL
# ============================================================

def validar_control_diario(
    control_id,
    validador_id,
    observacion=None
):
    """
    Marca un control como VALIDADO.
    """

    validador = obtener_usuario(
        validador_id
    )

    if not validar_rol_admin(
        validador
    ):

        raise PermissionError(
            "No tiene permisos para validar controles"
        )

    control = obtener_control_diario(
        control_id
    )

    if not control:

        raise ValueError(
            "Control no encontrado"
        )

    if control.estado == "VALIDADO":

        raise ValueError(
            "El control ya está validado"
        )

    control.estado = "VALIDADO"

    control.validado_por = validador.id

    control.validado_at = datetime.now(
        BOGOTA
    )

    control.observacion_validacion = (
        observacion
    )

    db.session.commit()

    return control


# ============================================================
# RECHAZAR CONTROL
# ============================================================

def rechazar_control_diario(
    control_id,
    validador_id,
    observacion
):
    """
    Marca un control como RECHAZADO.
    """

    validador = obtener_usuario(
        validador_id
    )

    if not validar_rol_admin(
        validador
    ):

        raise PermissionError(
            "No tiene permisos para rechazar controles"
        )

    if not observacion:

        raise ValueError(
            "Debe indicar el motivo del rechazo"
        )

    control = obtener_control_diario(
        control_id
    )

    if not control:

        raise ValueError(
            "Control no encontrado"
        )

    control.estado = "RECHAZADO"

    control.validado_por = validador.id

    control.validado_at = datetime.now(
        BOGOTA
    )

    control.observacion_validacion = (
        observacion
    )

    db.session.commit()

    return control



# ============================================================
# INTELLIFEET
# GENERADOR DE EXCEL CONSOLIDADO — CONTROL DIARIO
# ============================================================

import os

from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage

from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.page import PageMargins


# ============================================================
# GENERAR EXCEL CONSOLIDADO
# ============================================================
# ============================================================
# INTELLIFEET
# GENERADOR DE EXCEL CONSOLIDADO — CONTROL DIARIO
# ============================================================

def generar_excel_controles_diarios(
    controles,
    base_url=None,
    periodo_texto=None
):
    """
    Genera un Excel corporativo consolidado de Control Diario.

    HOJA 1:
    - Resumen ejecutivo profesional.
    - Logo IntelliFeet.
    - Información del reporte.
    - Indicadores KPI.
    - Distribución del reporte.

    HOJA 2:
    - Detalle consolidado.
    - Fotografías insertadas.
    - PDFs/archivos como enlaces.
    - Filtros.
    - Congelación de encabezados.
    - Formato corporativo.
    """

    wb = Workbook()

    # ========================================================
    # PALETA CORPORATIVA INTELLIFEET
    # ========================================================

    azul_oscuro = "172033"
    azul_navy = "1F2937"
    azul = "2563EB"
    azul_claro = "EFF6FF"

    dorado = "A47B20"
    dorado_claro = "F7F1E3"

    blanco = "FFFFFF"

    fondo = "F5F7FA"
    gris = "E5E7EB"
    gris_borde = "D9DEE7"
    gris_texto = "64748B"

    verde = "059669"
    verde_claro = "ECFDF5"

    rojo = "DC2626"
    rojo_claro = "FEF2F2"

    naranja = "D97706"
    naranja_claro = "FFF7ED"

    # ========================================================
    # BORDES
    # ========================================================

    borde_fino = Border(
        left=Side(
            style="thin",
            color=gris_borde
        ),
        right=Side(
            style="thin",
            color=gris_borde
        ),
        top=Side(
            style="thin",
            color=gris_borde
        ),
        bottom=Side(
            style="thin",
            color=gris_borde
        )
    )

    # ========================================================
    # FUENTES
    # ========================================================

    fuente_titulo = Font(
        name="Aptos Display",
        bold=True,
        size=20,
        color=blanco
    )

    fuente_header = Font(
        name="Aptos",
        bold=True,
        size=10,
        color=blanco
    )

    fuente_normal = Font(
        name="Aptos",
        size=10,
        color=azul_navy
    )

    fuente_gris = Font(
        name="Aptos",
        size=9,
        color=gris_texto
    )

    fuente_kpi = Font(
        name="Aptos Display",
        bold=True,
        size=20,
        color=azul_oscuro
    )

    fuente_kpi_label = Font(
        name="Aptos",
        bold=True,
        size=8,
        color=gris_texto
    )

    # ========================================================
    # ALINEACIONES
    # ========================================================

    alineacion_centro = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    alineacion_izquierda = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True
    )

    # ========================================================
    # INFORMACIÓN GENERAL
    # ========================================================

    # IMPORTANTE:
    # Se conserva exactamente la lógica que ya tenías.

    total = len(controles)

    validados = sum(
        1
        for c in controles
        if c.estado == "VALIDADO"
    )

    pendientes = sum(
        1
        for c in controles
        if c.estado == "PENDIENTE"
    )

    rechazados = sum(
        1
        for c in controles
        if c.estado == "RECHAZADO"
    )

    vehiculos = sum(
        1
        for c in controles
        if c.vehiculo_id is not None
    )

    maquinarias = sum(
        1
        for c in controles
        if c.maquinaria_id is not None
    )

    operadores = len({
        c.usuario_id
        for c in controles
        if c.usuario_id
    })

    # ========================================================
    # PERÍODO
    # ========================================================

    if periodo_texto:

        periodo = periodo_texto

    elif controles:

        fechas = [
            c.fecha
            for c in controles
            if c.fecha
        ]

        if fechas:

            fecha_min = min(fechas)
            fecha_max = max(fechas)

            periodo = (
                f"{fecha_min.strftime('%d/%m/%Y')} "
                f"al "
                f"{fecha_max.strftime('%d/%m/%Y')}"
            )

        else:

            periodo = "Período no especificado"

    else:

        periodo = "Sin registros"

    fecha_generacion = datetime.now(
        BOGOTA
    ).strftime(
        "%d/%m/%Y %H:%M"
    )

    # ========================================================
    # ========================================================
    # HOJA 1 — RESUMEN
    # ========================================================
    # ========================================================

    ws_resumen = wb.active

    ws_resumen.title = "Resumen"

    ws_resumen.sheet_view.showGridLines = False

    # ========================================================
    # ANCHOS
    # ========================================================

    ws_resumen.column_dimensions["A"].width = 3
    ws_resumen.column_dimensions["B"].width = 20
    ws_resumen.column_dimensions["C"].width = 21
    ws_resumen.column_dimensions["D"].width = 21
    ws_resumen.column_dimensions["E"].width = 21
    ws_resumen.column_dimensions["F"].width = 3

    # ========================================================
    # FONDO GENERAL
    # ========================================================

    for fila in range(1, 30):

        for columna in range(1, 7):

            ws_resumen.cell(
                fila,
                columna
            ).fill = PatternFill(
                "solid",
                fgColor=fondo
            )

    # ========================================================
    # CABECERA CORPORATIVA
    # ========================================================

    for fila in range(1, 4):

        for columna in range(1, 7):

            ws_resumen.cell(
                fila,
                columna
            ).fill = PatternFill(
                "solid",
                fgColor=azul_oscuro
            )

    ws_resumen.row_dimensions[1].height = 48
    ws_resumen.row_dimensions[2].height = 27
    ws_resumen.row_dimensions[3].height = 24
    ws_resumen.row_dimensions[4].height = 8

    # ========================================================
    # LOGO
    # ========================================================

    try:

        # Busca static/intellifeet.png partiendo desde
        # el directorio donde está este servicio.

        base_dir = Path(
            __file__
        ).resolve().parent

        posibles_rutas_logo = [

            base_dir / "static" / "intellifeet.png",

            base_dir.parent / "static" / "intellifeet.png",

            Path("static") / "intellifeet.png"

        ]

        ruta_logo = None

        for ruta in posibles_rutas_logo:

            if ruta.exists():

                ruta_logo = ruta

                break

        if ruta_logo:

            logo = ExcelImage(
                str(ruta_logo)
            )

            max_width = 125
            max_height = 42

            if logo.width and logo.height:

                escala = min(
                    max_width / logo.width,
                    max_height / logo.height
                )

                logo.width = int(
                    logo.width * escala
                )

                logo.height = int(
                    logo.height * escala
                )

            ws_resumen.add_image(
                logo,
                "B1"
            )

    except Exception as e:

        print(
            "Advertencia: no se pudo cargar "
            "el logo IntelliFeet:",
            e
        )

    # ========================================================
    # NOMBRE DEL SISTEMA
    # ========================================================

    ws_resumen.merge_cells(
        "C1:E1"
    )

    ws_resumen["C1"] = "INTELLIFEET"

    ws_resumen["C1"].font = fuente_titulo

    ws_resumen["C1"].alignment = Alignment(
        horizontal="right",
        vertical="center"
    )

    # ========================================================
    # TÍTULO
    # ========================================================

    ws_resumen.merge_cells(
        "C2:E2"
    )

    ws_resumen["C2"] = (
        "CONTROL DIARIO DE EVIDENCIAS OPERACIONALES"
    )

    ws_resumen["C2"].font = Font(
        name="Aptos",
        bold=True,
        size=11,
        color=blanco
    )

    ws_resumen["C2"].alignment = Alignment(
        horizontal="right",
        vertical="center"
    )

    # ========================================================
    # SUBTÍTULO
    # ========================================================

    ws_resumen.merge_cells(
        "C3:E3"
    )

    ws_resumen["C3"] = (
        "Reporte consolidado de gestión, "
        "validación y trazabilidad operacional"
    )

    ws_resumen["C3"].font = Font(
        name="Aptos",
        size=8,
        italic=True,
        color="CBD5E1"
    )

    ws_resumen["C3"].alignment = Alignment(
        horizontal="right",
        vertical="center"
    )

    # ========================================================
    # INFORME DE GESTIÓN OPERACIONAL
    # ========================================================

    ws_resumen.merge_cells(
        "B5:E5"
    )

    ws_resumen["B5"] = (
        "INFORME DE GESTIÓN OPERACIONAL"
    )

    ws_resumen["B5"].fill = PatternFill(
        "solid",
        fgColor=azul_navy
    )

    ws_resumen["B5"].font = Font(
        name="Aptos",
        bold=True,
        size=10,
        color=blanco
    )

    ws_resumen["B5"].alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    ws_resumen.row_dimensions[5].height = 26

    # ========================================================
    # INFORMACIÓN DEL DOCUMENTO
    # ========================================================

    informacion = [

        (
            "PERÍODO DEL REPORTE",
            periodo,
            "FECHA DE GENERACIÓN",
            fecha_generacion
        ),

        (
            "SISTEMA",
            "IntelliFeet",
            "MÓDULO",
            "Control Diario"
        ),

        (
            "TIPO DE DOCUMENTO",
            "Reporte Consolidado",
            "REGISTROS CONSOLIDADOS",
            total
        )

    ]

    fila_info = 6

    for (
        etiqueta_1,
        valor_1,
        etiqueta_2,
        valor_2
    ) in informacion:

        # ----------------------------------------------------
        # COLUMNA B
        # ----------------------------------------------------

        celda = ws_resumen.cell(
            fila_info,
            2,
            etiqueta_1
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=azul_claro
        )

        celda.font = Font(
            name="Aptos",
            bold=True,
            size=8,
            color=azul
        )

        celda.alignment = alineacion_izquierda

        celda.border = borde_fino

        # ----------------------------------------------------
        # COLUMNA C
        # ----------------------------------------------------

        celda = ws_resumen.cell(
            fila_info,
            3,
            valor_1
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=blanco
        )

        celda.font = Font(
            name="Aptos",
            bold=True,
            size=9,
            color=azul_navy
        )

        celda.alignment = alineacion_izquierda

        celda.border = borde_fino

        # ----------------------------------------------------
        # COLUMNA D
        # ----------------------------------------------------

        celda = ws_resumen.cell(
            fila_info,
            4,
            etiqueta_2
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=azul_claro
        )

        celda.font = Font(
            name="Aptos",
            bold=True,
            size=8,
            color=azul
        )

        celda.alignment = alineacion_izquierda

        celda.border = borde_fino

        # ----------------------------------------------------
        # COLUMNA E
        # ----------------------------------------------------

        celda = ws_resumen.cell(
            fila_info,
            5,
            valor_2
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=blanco
        )

        celda.font = Font(
            name="Aptos",
            bold=True,
            size=9,
            color=azul_navy
        )

        celda.alignment = alineacion_izquierda

        celda.border = borde_fino

        fila_info += 1

    # ========================================================
    # INDICADORES
    # ========================================================

    ws_resumen.merge_cells(
        "B10:E10"
    )

    ws_resumen["B10"] = (
        "INDICADORES DE CONTROL OPERACIONAL"
    )

    ws_resumen["B10"].fill = PatternFill(
        "solid",
        fgColor=azul_navy
    )

    ws_resumen["B10"].font = Font(
        name="Aptos",
        bold=True,
        size=10,
        color=blanco
    )

    ws_resumen["B10"].alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    ws_resumen.row_dimensions[10].height = 26

    # ========================================================
    # CUATRO KPI
    # ========================================================

    kpis = [

        (
            2,
            total,
            "CONTROLES TOTALES",
            azul_claro,
            azul_oscuro
        ),

        (
            3,
            validados,
            "VALIDADOS",
            verde_claro,
            verde
        ),

        (
            4,
            pendientes,
            "PENDIENTES",
            naranja_claro,
            naranja
        ),

        (
            5,
            rechazados,
            "RECHAZADOS",
            rojo_claro,
            rojo
        )

    ]

    for (
        columna,
        valor,
        etiqueta,
        color_fondo,
        color_texto
    ) in kpis:

        # ----------------------------------------------------
        # VALOR
        # ----------------------------------------------------

        celda_valor = ws_resumen.cell(
            11,
            columna,
            valor
        )

        celda_valor.fill = PatternFill(
            "solid",
            fgColor=color_fondo
        )

        celda_valor.font = Font(
            name="Aptos Display",
            bold=True,
            size=19,
            color=color_texto
        )

        celda_valor.alignment = alineacion_centro

        celda_valor.border = borde_fino

        # ----------------------------------------------------
        # ETIQUETA
        # ----------------------------------------------------

        celda_etiqueta = ws_resumen.cell(
            12,
            columna,
            etiqueta
        )

        celda_etiqueta.fill = PatternFill(
            "solid",
            fgColor=color_fondo
        )

        celda_etiqueta.font = Font(
            name="Aptos",
            bold=True,
            size=7,
            color=color_texto
        )

        celda_etiqueta.alignment = alineacion_centro

        celda_etiqueta.border = borde_fino

    ws_resumen.row_dimensions[11].height = 38
    ws_resumen.row_dimensions[12].height = 23

    # ========================================================
    # DISTRIBUCIÓN DEL REPORTE
    # ========================================================

    ws_resumen.merge_cells(
        "B14:E14"
    )

    ws_resumen["B14"] = (
        "DISTRIBUCIÓN DEL REPORTE"
    )

    ws_resumen["B14"].font = Font(
        name="Aptos",
        bold=True,
        size=10,
        color=blanco
    )

    ws_resumen["B14"].fill = PatternFill(
        "solid",
        fgColor=dorado
    )

    ws_resumen["B14"].alignment = alineacion_centro

    ws_resumen.row_dimensions[14].height = 24

    # ========================================================
    # DATOS DE DISTRIBUCIÓN
    # ========================================================

    distribucion = [

        (
            "Validados",
            validados,
            verde
        ),

        (
            "Pendientes",
            pendientes,
            naranja
        ),

        (
            "Rechazados",
            rechazados,
            rojo
        ),

        (
            "Vehículos",
            vehiculos,
            azul
        ),

        (
            "Maquinaria",
            maquinarias,
            dorado
        ),

        (
            "Operadores",
            operadores,
            azul_navy
        )

    ]

    fila = 15

    for (
        etiqueta,
        valor,
        color
    ) in distribucion:

        # ----------------------------------------------------
        # ETIQUETA
        # ----------------------------------------------------

        celda = ws_resumen.cell(
            fila,
            2,
            etiqueta
        )

        celda.font = Font(
            name="Aptos",
            bold=True,
            size=9,
            color=azul_navy
        )

        celda.alignment = alineacion_izquierda

        celda.fill = PatternFill(
            "solid",
            fgColor=blanco
        )

        celda.border = borde_fino

        # ----------------------------------------------------
        # VALOR
        # ----------------------------------------------------

        celda = ws_resumen.cell(
            fila,
            3,
            valor
        )

        celda.font = Font(
            name="Aptos",
            bold=True,
            size=11,
            color=color
        )

        celda.alignment = alineacion_centro

        celda.fill = PatternFill(
            "solid",
            fgColor=blanco
        )

        celda.border = borde_fino

        # ----------------------------------------------------
        # CELDAS RESTANTES
        # ----------------------------------------------------

        for columna in (4, 5):

            celda = ws_resumen.cell(
                fila,
                columna,
                ""
            )

            celda.fill = PatternFill(
                "solid",
                fgColor=blanco
            )

            celda.border = borde_fino

        fila += 1

    # ========================================================
    # PIE DEL RESUMEN
    # ========================================================

    ws_resumen.merge_cells(
        "B23:E23"
    )

    ws_resumen["B23"] = (
        "INTELLIFEET · Gestión y trazabilidad operacional"
    )

    ws_resumen["B23"].font = Font(
        name="Aptos",
        italic=True,
        size=8,
        color=gris_texto
    )

    ws_resumen["B23"].alignment = alineacion_centro

    # ========================================================
    # IMPRESIÓN RESUMEN
    # ========================================================

    ws_resumen.page_setup.orientation = "portrait"

    ws_resumen.page_setup.paperSize = (
        ws_resumen.PAPERSIZE_A4
    )

    ws_resumen.page_setup.fitToWidth = 1

    ws_resumen.page_setup.fitToHeight = 1

    ws_resumen.sheet_properties.pageSetUpPr.fitToPage = True

    ws_resumen.page_margins = PageMargins(
        left=0.35,
        right=0.35,
        top=0.50,
        bottom=0.50,
        header=0.20,
        footer=0.20
    )

    ws_resumen.oddFooter.center.text = (
        "INTELLIFEET · Reporte Consolidado"
    )

    ws_resumen.oddFooter.center.size = 8

    ws_resumen.oddFooter.right.text = (
        "Página &[Page] de &[Pages]"
    )

    ws_resumen.oddFooter.right.size = 8

    # ========================================================
    # ========================================================
    # HOJA 2 — DETALLE
    # ========================================================
    # ========================================================

    ws = wb.create_sheet(
        "Detalle"
    )

    ws.sheet_view.showGridLines = False

    # ========================================================
    # COLUMNAS
    # ========================================================

    columnas = [

        "FECHA",

        "ACTIVO",

        "TIPO ACTIVO",

        "PLACA / CÓDIGO",

        "OPERADOR",

        "EVIDENCIA",

        "TIPO ARCHIVO",

        "ESTADO",

        "OBSERVACIONES",

        "VALIDADO POR",

        "FECHA VALIDACIÓN",

        "OBSERVACIÓN VALIDACIÓN",

        "FECHA CARGA"

    ]

    # ========================================================
    # ENCABEZADO
    # ========================================================

    for columna, titulo in enumerate(
        columnas,
        start=1
    ):

        celda = ws.cell(
            row=1,
            column=columna,
            value=titulo
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=azul_oscuro
        )

        celda.font = fuente_header

        celda.alignment = alineacion_centro

        celda.border = borde_fino

    ws.row_dimensions[1].height = 30

    # ========================================================
    # ANCHOS
    # ========================================================

    anchos = {

        1: 14,

        2: 22,

        3: 18,

        4: 20,

        5: 28,

        6: 32,

        7: 16,

        8: 16,

        9: 40,

        10: 28,

        11: 22,

        12: 40,

        13: 22

    }

    for columna, ancho in anchos.items():

        ws.column_dimensions[
            get_column_letter(columna)
        ].width = ancho

    # ========================================================
    # DATOS
    # ========================================================

    fila = 2

    for control in controles:

        # ----------------------------------------------------
        # ACTIVO
        # ----------------------------------------------------

        if control.vehiculo:

            tipo_activo = "VEHÍCULO"

            activo = (
                control.vehiculo.placa
                or ""
            )

            placa_codigo = (
                control.vehiculo.placa
                or ""
            )

        elif control.maquinaria:

            tipo_activo = "MAQUINARIA"

            activo = (
                control.maquinaria.codigo
                or ""
            )

            placa_codigo = (
                control.maquinaria.codigo
                or ""
            )

        else:

            tipo_activo = ""

            activo = ""

            placa_codigo = ""

        # ----------------------------------------------------
        # USUARIOS
        # ----------------------------------------------------

        operador = (
            control.usuario.nombre
            if control.usuario
            else ""
        )

        validador = (
            control.validador.nombre
            if control.validador
            else ""
        )

        # ----------------------------------------------------
        # FECHAS
        # ----------------------------------------------------

        fecha_control = (
            control.fecha.strftime(
                "%d/%m/%Y"
            )
            if control.fecha
            else ""
        )

        fecha_validacion = (
            control.validado_at.strftime(
                "%d/%m/%Y %H:%M"
            )
            if control.validado_at
            else ""
        )

        fecha_carga = (
            control.created_at.strftime(
                "%d/%m/%Y %H:%M"
            )
            if control.created_at
            else ""
        )

        # ----------------------------------------------------
        # DATOS
        # ----------------------------------------------------

        datos = [

            fecha_control,

            activo,

            tipo_activo,

            placa_codigo,

            operador,

            control.archivo_nombre or "",

            (
                control.archivo_tipo.upper()
                if control.archivo_tipo
                else ""
            ),

            control.estado or "",

            control.observaciones or "",

            validador,

            fecha_validacion,

            control.observacion_validacion or "",

            fecha_carga

        ]

        # ----------------------------------------------------
        # INSERTAR FILA
        # ----------------------------------------------------

        for columna, valor in enumerate(
            datos,
            start=1
        ):

            celda = ws.cell(
                row=fila,
                column=columna,
                value=valor
            )

            celda.border = borde_fino

            celda.alignment = alineacion_izquierda

            celda.font = fuente_normal

            if fila % 2 == 0:

                celda.fill = PatternFill(
                    "solid",
                    fgColor="FAFBFC"
                )

        # ====================================================
        # ESTADO
        # ====================================================

        celda_estado = ws.cell(
            row=fila,
            column=8
        )

        celda_estado.alignment = alineacion_centro

        if control.estado == "VALIDADO":

            celda_estado.fill = PatternFill(
                "solid",
                fgColor=verde
            )

            celda_estado.font = Font(
                name="Aptos",
                bold=True,
                color=blanco,
                size=9
            )

        elif control.estado == "RECHAZADO":

            celda_estado.fill = PatternFill(
                "solid",
                fgColor=rojo
            )

            celda_estado.font = Font(
                name="Aptos",
                bold=True,
                color=blanco,
                size=9
            )

        elif control.estado == "PENDIENTE":

            celda_estado.fill = PatternFill(
                "solid",
                fgColor=naranja
            )

            celda_estado.font = Font(
                name="Aptos",
                bold=True,
                color=blanco,
                size=9
            )

        # ====================================================
        # TIPO DE ACTIVO
        # ====================================================

        celda_tipo = ws.cell(
            row=fila,
            column=3
        )

        celda_tipo.alignment = alineacion_centro

        if tipo_activo == "VEHÍCULO":

            celda_tipo.font = Font(
                name="Aptos",
                bold=True,
                color=azul
            )

        elif tipo_activo == "MAQUINARIA":

            celda_tipo.font = Font(
                name="Aptos",
                bold=True,
                color=dorado
            )

        # ====================================================
        # EVIDENCIA
        # ====================================================

        ruta_archivo = os.path.join(
            "uploads",
            control.archivo_path
        )

        extension = (
            control.archivo_tipo or ""
        ).lower().replace(".", "")

        extensiones_imagen = {
            "jpg",
            "jpeg",
            "png",
            "webp"
        }

        if (
            extension in extensiones_imagen
            and os.path.exists(ruta_archivo)
        ):

            try:

                imagen = ExcelImage(
                    ruta_archivo
                )

                max_width = 250
                max_height = 170

                ancho_original = imagen.width
                alto_original = imagen.height

                if (
                    ancho_original
                    and alto_original
                ):

                    escala = min(
                        max_width / ancho_original,
                        max_height / alto_original,
                        1
                    )

                    imagen.width = int(
                        ancho_original * escala
                    )

                    imagen.height = int(
                        alto_original * escala
                    )

                ws.add_image(
                    imagen,
                    f"F{fila}"
                )

                ws.row_dimensions[
                    fila
                ].height = 130

                celda_archivo = ws.cell(
                    fila,
                    6
                )

                celda_archivo.value = (
                    control.archivo_nombre
                    or "Evidencia"
                )

                celda_archivo.alignment = Alignment(
                    horizontal="center",
                    vertical="bottom",
                    wrap_text=True
                )

                celda_archivo.font = Font(
                    name="Aptos",
                    size=8,
                    color=gris_texto
                )

            except Exception:

                celda_archivo = ws.cell(
                    fila,
                    6,
                    control.archivo_nombre
                )

                celda_archivo.font = Font(
                    color=azul,
                    underline="single"
                )

        else:

            celda_archivo = ws.cell(
                fila,
                6,
                control.archivo_nombre
            )

            celda_archivo.alignment = alineacion_centro

            if base_url:

                url = (
                    base_url.rstrip("/")
                    + "/uploads/"
                    + control.archivo_path
                )

                celda_archivo.hyperlink = url

                celda_archivo.font = Font(
                    name="Aptos",
                    color="0563C1",
                    underline="single",
                    bold=True
                )

        fila += 1

    # ========================================================
    # TABLA EXCEL
    # ========================================================

    ultima_fila = max(
        fila - 1,
        1
    )

    if ultima_fila >= 2:

        tabla = Table(
            displayName="TablaControlDiario",
            ref=f"A1:M{ultima_fila}"
        )

        estilo_tabla = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False
        )

        tabla.tableStyleInfo = estilo_tabla

        ws.add_table(
            tabla
        )

    # ========================================================
    # FILTROS
    # ========================================================

    if ultima_fila >= 2:

        ws.auto_filter.ref = (
            f"A1:M{ultima_fila}"
        )

    # ========================================================
    # CONGELAR ENCABEZADO
    # ========================================================

    ws.freeze_panes = "A2"

    # ========================================================
    # CONFIGURACIÓN DE IMPRESIÓN
    # ========================================================

    ws.page_setup.orientation = "landscape"

    ws.page_setup.paperSize = (
        ws.PAPERSIZE_A4
    )

    ws.page_setup.fitToWidth = 1

    ws.page_setup.fitToHeight = 0

    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins = PageMargins(
        left=0.25,
        right=0.25,
        top=0.50,
        bottom=0.50,
        header=0.20,
        footer=0.20
    )

    ws.print_title_rows = "1:1"

    # ========================================================
    # PIE DE PÁGINA
    # ========================================================

    ws.oddFooter.center.text = (
        "INTELLIFEET · Control Diario"
    )

    ws.oddFooter.center.size = 8

    ws.oddFooter.right.text = (
        "Página &[Page] de &[Pages]"
    )

    ws.oddFooter.right.size = 8

    # ========================================================
    # PROPIEDADES DEL DOCUMENTO
    # ========================================================

    wb.properties.title = (
        "IntelliFeet - Control Diario"
    )

    wb.properties.subject = (
        "Reporte consolidado de controles diarios"
    )

    wb.properties.creator = (
        "IntelliFeet"
    )

    wb.properties.description = (
        "Reporte corporativo consolidado "
        "de evidencias de Control Diario."
    )

    wb.properties.keywords = (
        "IntelliFeet, Control Diario, "
        "Operaciones, Evidencias, Flota"
    )

    # ========================================================
    # GUARDAR EN MEMORIA
    # ========================================================

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    return output